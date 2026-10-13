"""Build the Chinese onboarding manual from verified product UI and contracts."""
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "织界产品设计与代码生成使用指南.docx"
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.7)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
sec.header_distance, sec.footer_distance = Cm(.7), Cm(.7)

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
    st = doc.styles[name]
    st.font.name = "Arial"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = Pt(15)
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Title"].font.size = Pt(25)
doc.styles["Title"].paragraph_format.line_spacing = Pt(31)
doc.styles["Subtitle"].font.size = Pt(11)
doc.styles["Subtitle"].font.italic = False
doc.styles["Heading 1"].font.size = Pt(17)
doc.styles["Heading 1"].paragraph_format.line_spacing = Pt(24)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 2"].font.size = Pt(12)
doc.styles["Heading 2"].paragraph_format.line_spacing = Pt(17)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
doc.styles["Caption"].font.size = Pt(8.5)
doc.styles["Caption"].paragraph_format.line_spacing = Pt(12)
doc.styles["Caption"].paragraph_format.space_after = Pt(7)
header = sec.header.paragraphs[0]
header.add_run("织界使用指南  |  产品设计与研发代码生成").font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("2026 年 10 月 14 日   ·   ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "织界产品设计与代码生成使用指南"
doc.core_properties.subject = "安装 账号注册 产品原型设计 AI 连接 React 代码生成"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""


def p(text, style=None, bold=False):
    para = doc.add_paragraph(style=style)
    para.add_run(text).bold = bold
    return para


def h(text):
    doc.add_heading(text, level=2)


def page(title):
    para = doc.add_heading(title, level=1)
    para.paragraph_format.page_break_before = True


def steps(items):
    for i, text in enumerate(items, 1):
        para = p(f"{i}. {text}")
        para.paragraph_format.left_indent = Cm(.2)
        para.paragraph_format.first_line_indent = Cm(-.2)


def bullets(items):
    for text in items:
        p("• " + text)


def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing = Pt(14)
    para.paragraph_format.left_indent = Cm(.25)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F2F4F7")
    para._p.get_or_add_pPr().append(shade)
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")


def link(label, url):
    para = doc.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "244E78")
    props.append(color)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    props.append(fonts)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label + "  " + url
    run.append(text)
    hyperlink.append(run)
    para._p.append(hyperlink)


def image(name, caption, width=16.5):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run()
    pic = run.add_picture(str(SHOTS / name), width=Cm(width))
    pic._inline.docPr.set("descr", caption)
    cap = p(caption, "Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        b = OxmlElement("w:" + side)
        for k, v in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            b.set(qn("w:" + k), v)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for j, row in enumerate(t.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EEF5" if j == 0 else ("F7F9FC" if j % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "bottom", "left", "right"]:
                m = OxmlElement("w:" + side)
                m.set(qn("w:w"), "85")
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            cell._tc.get_or_add_tcPr().append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
                para.paragraph_format.line_spacing = Pt(14)
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    run.bold = j == 0
    p("")


# Screenshots are cropped only to improve legibility or exclude personal data.
Image.open(SHOTS / "09-preview-drawer.jpg").crop((1160, 0, 1800, 986)).save(SHOTS / "09-preview-drawer-crop.png")

p("织界产品设计与代码生成\n使用指南", "Title")
p("从下载安装到可编辑原型与 React 项目代码", "Subtitle")
p("核对日期 2026 年 10 月 14 日  ·  Desktop 0.0.62  ·  CLI 0.0.84  ·  资源 0.0.87")
p("织界基于 Plasmic，支持在可视化编辑器中设计页面，再由连接桌面 MCP 的 AI 客户端读取设计、生成目标工程中的 React 与 TypeScript 代码。本指南面向首次使用的产品、设计与研发人员，按实际操作顺序介绍安装、账号、设计、评审和代码交付。")
h("选择你的使用路线")
table(["你的目标", "需要准备", "阅读顺序"], [
    ["手工设计和评审", "拾光账号；浏览器或 Desktop", "下载安装 → 账号 → 工作区 → 编辑器 → 预览与分享"],
    ["通过外部 AI 设计原型", "Desktop；可用 AI 客户端；MCP；CLI 与 Skill", "账号 → Desktop → MCP → CLI → AI 设计"],
    ["从设计生成研发代码", "设计访问权限；Desktop；AI 客户端；目标代码工程", "MCP → CLI → 设计交接 → 代码生成 → 验收"]
], [3.2, 6.5, 6.8])
h("全流程")
p("注册并登录 → 安装 Desktop → 连接 AI 客户端 → 安装 Skill → 创建设计 → 预览评审 → 交给研发 → 生成代码 → 在真实工程中验收。")
h("本手册内容")
p("02 下载安装　03 注册与网页登录　04 Desktop 登录\n05 MCP 连接　06 CLI 与 Skill　07 工作区与项目\n08 编辑器操作　09 外部 AI 设计　10 内置 AI 编辑\n11 预览验收　12 分享与交接　13 研发代码生成\n14 增量更新与验收　15 常见问题　16 入口与版本")
p("网页端与 Desktop 使用同一个账号和云端项目。外部 AI 通过 MCP 操作设计时，需要保持 Desktop 运行并登录。Desktop 的项目加载、保存仍需要联网。")

page("02 下载并安装应用")
link("官网下载入口", "https://shiguanglab.com/projects/zhijie#download")
image("01-download.jpg", "图 1  官网下载区域按操作系统提供安装包", 16.5)
table(["设备", "选择的安装包", "安装方式"], [
    ["Mac M 系列芯片", "macOS Apple Silicon · DMG", "打开 DMG，将 Plasmic 拖到 Applications"],
    ["Mac Intel 处理器", "macOS Intel · DMG", "打开 DMG，将 Plasmic 拖到 Applications"],
    ["Windows x64", "Windows · EXE", "运行安装程序，按向导完成安装"],
    ["Linux x64", "Linux · AppImage", "赋予文件可执行权限后运行"]
], [3.3, 5.3, 7.9])
p("Mac 可在苹果菜单的“关于本机”查看芯片类型；出现“芯片”与 M 系列名称时选择 Apple Silicon，Intel 设备选择 Intel。下载页自动识别设备，但仍应核对实际机型。")
p("Linux 示例命令中的文件名需要替换为实际下载文件名：")
code('chmod +x "Plasmic-版本号-linux-x86_64.AppImage"\n./Plasmic-版本号-linux-x86_64.AppImage')
p("只使用网页端时，直接进入 Studio，无需安装 Desktop。使用安装包也无需自行安装 Docker、数据库或本仓库源码；Node.js 仅在使用 CLI／Skill 时需要。若系统阻止打开安装包，确认下载来源后按操作系统的安全提示处理，企业管理设备可联系 IT。")

page("03 注册账号并登录网页")
link("从 Studio 进入统一账号登录", "https://studio.plasmic.shiguanglab.com/")
p("首次访问 Studio 会进入拾光统一账号页面。已有拾光账号可直接登录；没有账号时点击“去注册”。")
para = doc.add_paragraph()
para.alignment = WD_ALIGN_PARAGRAPH.CENTER
para.paragraph_format.keep_with_next = True
para.paragraph_format.line_spacing = 1.0
for name in ["03-register.jpg", "02-login.jpg"]:
    pic = para.add_run().add_picture(str(SHOTS / name), width=Cm(7.2))
    pic._inline.docPr.set("descr", "拾光账号注册与登录表单")
    para.add_run("  ")
p("图 2  左侧为注册表单，右侧为账号登录表单", "Caption")
h("创建账号")
steps([
    "填写用户名：3–20 个字符，可包含字母、数字和下划线。",
    "填写可接收邮件的邮箱；设置至少 8 位、同时包含字母和数字的密码，再填写确认密码。",
    "阅读页面的用户协议和隐私政策，自行确认同意后点击“创建账号”。",
    "若提示验证邮件已发送，在邮箱中完成验证后登录；若直接返回产品页面，继续下一步。"
])
h("选择正确的登录方式")
p("账号登录使用用户名和密码；邮箱登录按页面提示获取并输入邮箱验证码。页面还提供 GitHub、Google 和飞书入口，使用你注册或绑定时对应的身份。第三方身份与密码账号不能仅凭相同邮箱就视为同一账号。")
p("收不到邮件时检查垃圾邮件和邮箱拼写；用户名或邮箱已使用时改用已有账号登录。需要协助时点击“无法登录？”。")

page("04 登录 Desktop")
image("12-desktop-login.png", "图 3  Desktop 登录入口示例  2026 年 10 月 11 日界面记录", 16.5)
steps([
    "启动已安装的 Plasmic，点击登录按钮；Desktop 会打开系统浏览器。",
    "在拾光统一账号页面登录。新用户可从同一页面进入注册。",
    "浏览器完成授权后，允许打开 Plasmic Desktop；若没有自动返回，点击授权页面上的“Open Plasmic Desktop”或对应中文按钮。",
    "回到 Desktop，确认出现工作区或原项目。浏览器显示授权成功后，还需要以 Desktop 实际进入项目为登录完成依据。"
])
h("日常使用")
bullets([
    "保持 Desktop 运行；外部 AI 工具调用期间不要退出应用。",
    "网页登录成功不等于 Desktop 已建立登录会话，首次使用仍需完成应用授权。",
    "会话过期时重新登录，应用会保留待返回的页面位置。",
    "账号设置进入拾光账号中心；网页与 Desktop 的项目、工作区和权限一致。"
])
h("应用更新")
p("Desktop 启动后会检查更新并在后台下载。下载完成后通过软件更新窗口选择安装和重启；安装前保存当前设计。也可使用应用菜单的“检查更新”。应用更新、CLI 更新和 Skill／规范更新是三件不同的操作。")

page("05 将 AI 客户端连接到 Desktop")
p("MCP 是 AI 客户端连接本地编辑器的通道。先启动并登录 Desktop，再打开需要处理的项目。仅安装 Skill 不会自动完成 MCP 连接。")
steps([
    "在 Desktop 的应用菜单选择“AI → MCP”。Mac 菜单位于屏幕顶部，Windows／Linux 使用应用菜单。",
    "在“MCP 连接”窗口中，开启你实际使用的 AI 客户端开关。",
    "重启 AI 客户端，或刷新其 MCP 配置；检查是否出现 plasmic-desktop 服务。",
    "在新对话中要求 AI 读取应用状态，并确认项目名称、登录状态和工具可用性。"
])
image("11-mcp.png", "图 4  Desktop MCP 设置示例  2026 年 10 月 10 日界面记录", 12.4)
p("内置自动配置支持 Claude Code CLI、Codex CLI、Gemini CLI、Antigravity 2.0、OpenCode CLI、Kiro CLI 与 Claude Desktop。其他 MCP 客户端可点击“复制 MCP 配置”，粘贴到其本地 MCP 服务配置中。Skill 安装器支持的客户端范围与这里的 MCP 自动配置列表不同。")
p("移动 Desktop 安装位置后，应重新生成或复制配置，因为配置使用应用的绝对路径。发生配置冲突时先检查提示，不覆盖其他服务。可用的本地 AI 客户端需由使用者自行安装并配置模型。")

page("06 安装 CLI 和 Plasmic Skill")
p("CLI 的名称是 plasmickit，npm 包名是 @plasmickit/cli。它负责安装 Skill、更新与校验工作流规范；实际设计读取、原型编辑和代码实现由 AI 客户端配合 Desktop MCP 完成。")
h("准备 Node.js")
link("Node.js 官方下载页", "https://nodejs.org/en/download")
p("使用 Node.js 22.12 或更高版本。安装后重新打开终端，在 macOS／Linux 终端或 Windows PowerShell 中执行：")
code("node --version\nnpm --version")
h("安装并核对版本")
code("npm install -g @plasmickit/cli@latest\nplasmickit --version\nplasmickit skill install\nplasmickit version check")
p("skill install 会检测本机支持的 AI 客户端，并安装同一份 Plasmic Skill。检查输出中的安装结果；如果未检测到客户端，先安装或初始化目标客户端，再重新执行。安装完成后启动一个新对话。")
h("不进行全局安装的方式")
code("npx -y @plasmickit/cli@latest skill install\nnpx -y @plasmickit/cli@latest context resolve --mode prototype")
h("加载设计与代码生成规范")
code("plasmickit context resolve --mode prototype\nplasmickit context resolve --mode codegen")
p("每次 resolve 都检查线上资源并校验缓存。AI 应按返回结果读取 mustRead 文件；正常使用可让 AI 调用 Plasmic Skill 来完成该步骤。模式从设计切换到研发时，重新加载 codegen 规范。下载或校验失败时先解决网络问题，再继续依赖该规范的任务。")
h("后续更新")
code("npm install -g @plasmickit/cli@latest\nplasmickit skill install\nplasmickit references update")
p("以上命令分别更新 CLI、Skill 入口和规范资源。版本号不同是正常情况；无需为了资源更新重新安装 Desktop。")

page("07 选择工作区并创建项目")
image("04-projects-crop.png", "图 5  项目列表支持搜索与工作区筛选", 16.5)
steps([
    "登录后进入“项目”或“个人工作区”。团队协作时先选择对应组织与工作区；看不到目标工作区时请管理员邀请。",
    "点击“新建项目”，填写项目名称，并选择项目所属工作区。",
    "从空白项目或当前可用模板开始，点击“创建”。本指南截图只展示创建窗口。",
    "点击项目卡片进入编辑器。已有项目优先直接打开，避免为同一需求重复建项目。"
])
image("05-create-project-crop.png", "图 6  新建项目需要名称与工作区", 5.5)
p("建议项目名称包含需求编号或业务模块，例如“REQ075 客群管理”。个人账号注册不会自动授予既有团队项目的权限。模板是否可见取决于工作区配置和素材库访问权限。")

page("08 在编辑器中设计页面")
image("06-editor.jpg", "图 7  文档与图层在左侧，画布居中，页面数据与属性在右侧", 16.5)
h("建立页面结构")
steps([
    "点击左侧“图层”，切换到“文档”，使用“新建”创建页面、组件或画板组。需要可访问地址的业务页面应创建 Page，并设置 URL 路径。",
    "通过“插入”放置布局与组件；通过“库”使用项目可访问的组件库。管理后台优先复用现有 AppShell、表单、表格和业务组件。",
    "在图层树中选择节点，在右侧修改组件属性、插槽内容与样式。业务文案可用中文；编辑节点名称使用有意义的英文 PascalCase，例如 CustomerTable。",
    "配置事件、状态、查询和页面导航。视觉上的按钮并不自动具备查询或保存业务数据的行为。",
    "查看顶部“已保存”状态，再刷新或重新打开页面，确认设计保留。"
])
h("常用控件")
table(["入口", "用途"], [
    ["图层与文档", "定位页面、组件和嵌套节点；选择实际编辑对象"],
    ["插入与库", "添加节点，复用已有组件和素材"],
    ["右侧属性与数据", "修改真实组件 Props、Slot、页面路由、状态与查询"],
    ["缩放与平移", "调整观察范围；“缩放以显示全部”帮助找回画布"],
    ["撤销与重做", "恢复最近的编辑操作"]
], [4.5, 12])
page("09 使用外部 AI 生成可编辑原型")
p("适合从需求文档创建多页面原型，或对现有项目进行有范围的调整。开始前完成 Desktop 登录、MCP 连接和 Skill 安装，并在 AI 客户端提供可读取的需求文档。")
h("可直接复制的设计指令")
code("使用 Plasmic skill 为 REQ075 设计客群管理原型。\n组织：<组织名称>；工作区：<工作区名称>。\n项目：<已有项目名称或明确要求新建>。\n需求文档：<文件路径或可访问链接>。\n\n先读取需求，确认页面清单、字段、状态与交互；\n检查当前组件库与模板的真实契约，再设计可编辑页面。\n范围：桌面端，视口 1440 × 900，业务文案使用简体中文。\n包含列表查询、重置、分页、详情和新增表单的实际交互。\n未接入真实接口时明确说明模拟数据范围。\n在 Preview 中验证，再保存、重新打开并读回确认。\n交付项目链接、页面路由、已验证行为和待确认项。")
h("怎样提供有效输入")
bullets([
    "说明业务目标和使用角色，并提供完整需求文件，而非仅一张页面截图。",
    "列出页面、字段含义、枚举、必填规则、动作的成功／失败状态及页面跳转。",
    "指定组织、工作区、项目与访问范围；重名项目应附项目链接。",
    "可附视觉参考图或 design.pen，但交互和数据规则仍由需求定义。",
    "修改现有设计时写清受影响页面与区域，避免让 AI 无范围地重做项目。"
])
h("局部调整指令")
code("使用 Plasmic skill 修改 <项目链接> 中的 CustomerList 页面。\n仅调整筛选区：增加状态筛选，并在重置时清空已应用条件、回到第一页。\n保留表格、路由和其他页面。验证查询与重置，再保存并重新读回。")
p("AI 执行期间查看任务进度和当前项目。任务结束后仍需按需求验收；“已生成”或模型校验通过不能替代 Preview 中的真实操作验证。")
p("独立后台页面通常保留一个 AppShell；接入已有应用时复用宿主外壳。表格列、表单校验和弹层开关要按组件真实契约配置。")

page("10 使用内置 AI 进行页面编辑")
image("07-ai.jpg", "图 8  编辑器中的 AI 助手可对当前选区或页面提出修改", 16.5)
p("内置 AI 助手用于当前选区或页面的编辑，是另一种入口。外部 AI 客户端的完整 Skill 工作流仍按前述 MCP 与 CLI 步骤使用。")
steps([
    "在画布或图层树中先选择目标区域；没有选区时核对页面名称。",
    "点击左侧“AI”，在输入框中描述具体修改，可添加需求文件或参考图片。",
    "选择界面中实际可用的模型，确认编辑对象后发送请求。任务执行期间编辑范围保持固定。",
    "查看进度、结果和失败提示，检查改动；不符合需求时使用撤销或提出有范围的修正。",
    "在 Preview 中检查交互，再确认保存状态。"
])
h("局部编辑示例")
code("将当前筛选区改为两行布局，保留现有字段和事件。\n统一标签宽度，查询和重置按钮放在第二行右侧。\n不要修改表格列、路由或其他页面。")
p("当前线上环境已启用局部 AI 编辑；新页面生成功能是否可用以当前界面与服务配置为准。本次核对的环境未启用内置 AI 新页面生成。从需求建立完整项目，可使用第 09 节的外部 AI 与 Skill 流程。图中没有实际发送 AI 请求。")

page("11 预览并验证交互")
image("08-preview.jpg", "图 9  从顶部预览入口查看可运行的示例页面", 16.5)
steps([
    "点击顶部“预览”或“在新标签页中打开预览”，进入可交互页面。",
    "检查筛选、重置、分页、Tab、路由与弹层；按实际需求测试空数据、失败和只读状态。",
    "打开新增或编辑表单，检查必填提示、取消、再次打开与提交后的行为。",
    "回到编辑器修正问题；确认已保存后重新打开并复查。"
])
image("09-preview-drawer-crop.png", "图 10  新增资源抽屉的运行界面", 5.6)
p("Admin Templates 仅用于操作示例，展示数据为虚构内容。模板的业务查询、保存和项目路由仍需由使用模板的项目绑定；图中打开抽屉不代表已经接入生产接口。")

page("12 分享设计并交给研发")
image("10-share.jpg", "图 11  项目菜单中的分享窗口提供链接和邀请入口", 16)
steps([
    "先确认顶部显示“已保存”，记录本次评审的页面与范围。",
    "打开顶部项目菜单，选择“分享项目”。复制链接供评审，或按团队流程邀请指定成员。",
    "根据协作需要选择角色。研发读取设计需要相应访问权限；需要修改设计时再授予合适的编辑权限。",
    "请接收者登录自己的拾光账号，确认可以打开设计与预览。复制链接不会自动改变访问权限。",
    "将需求、页面链接、验收结论与接口说明一起交接，明确待确认事项。"
])
h("研发交接清单")
bullets([
    "项目链接、目标页面与路由、确认过的设计版本或 revision。",
    "需求编号、字段含义、枚举、校验、页面状态和交互规则。",
    "目标代码工程、公共业务组件、组件库版本和样式规范。",
    "真实接口文档、参数与响应、鉴权、权限和错误处理规则。",
    "图片与图标来源，哪些内容为模拟数据，哪些行为尚未联调。"
])
h("保存  发布  部署")
p("保存保留编辑模型；发布建立设计的发布版本或执行界面中指定的发布流程；部署将研发应用上线。三者分别处理。通过 AI 生成代码后，仍应按目标工程的代码审查、构建与上线流程交付；本手册的示例没有发布项目或修改共享权限。")

page("13 研发从设计生成项目代码")
p("研发在目标代码工程中启动 AI 客户端，连接已经登录的 Desktop，提供设计项目链接与页面范围。AI 读取真实页面模型、组件属性、插槽、状态与事件，然后按目标工程实现原生 React／TypeScript。")
steps([
    "准备已拉取的目标工程和设计访问权限，确认现有 package.json、路由、公共组件、请求与权限封装。",
    "在 Desktop 打开目标项目；让 AI 核对当前项目名称和目标 Page，读取隐藏 Tab、关闭弹层等实际设计分支。",
    "加载 codegen 工作流，提供真实接口契约。接口缺失时明确当前任务是静态演示，或列出待联调部分。",
    "要求 AI 将实现接入真实路由，复用工程已有组件与样式方案，再运行相关检查。"
])
code("plasmickit context resolve --mode codegen")
h("可直接复制的研发指令")
code("使用 Plasmic skill，读取 <设计项目链接> 的 CustomerList 页面。\n目标工程：<代码工程路径>；需求：<需求文档路径>。\n目标路由：/customers；接口文档：<接口路径或链接>。\n\n先读取页面模型、组件契约、状态、事件和隐藏分支；\n核对本工程实际依赖、公共组件、路由、请求与权限封装。\n在现有架构中实现 React/TypeScript，复用宿主外壳。\n保留查询、重置、分页、详情、表单校验和提交流程。\n不要假设原型中的模拟数据就是已有生产接口。\n完成后执行工程已有的相关类型检查、测试与构建，\n从真实应用入口核对交互和视觉，说明未验证项。\n交付变更文件、路由、数据边界和实际验证结果。")
h("明确生成方式")
p("本指南采用 Skill 驱动的原生代码实现流程。plasmickit 本身不输出页面代码；截图用于视觉核对，组件与交互以真实模型及实现为依据。该流程不调用 export_code，也不要求安装官方 @plasmicapp/cli、Loader 或同步配置。若团队明确采用官方同步方案，应单独按该方案的契约处理。")
p("在已有工程中沿用实际技术栈，不为单页生成升级组件库。新建独立项目且未指定技术栈时，当前 codegen 规范的默认组合是 React、TypeScript、Ant Design 6 和 Vite。")

page("14 增量更新与代码验收")
h("设计调整后如何更新代码")
p("设计变更先在编辑器保存并确认。研发让 AI 重新读取受影响页面与子树，对照现有代码更新承担该职责的文件；保留与本次改动无关的页面、接口和业务逻辑。")
code("使用 Plasmic skill 更新 <目标工程路径> 的 CustomerList。\n设计源：<项目链接>；仅同步状态筛选及查询重置行为。\n重新读取相关模型与接口，保留其他布局、路由和业务逻辑。\n运行受影响的检查，并在真实入口验证。")
h("验收重点")
table(["检查范围", "完成依据"], [
    ["编译与类型", "执行工程真实存在的类型检查和构建命令，记录结果"],
    ["页面入口", "真实路由可进入；Provider、主题与宿主集成正确"],
    ["查询和分页", "参数正确；查询、重置回正确页；total 与数据一致"],
    ["表单与弹层", "必填和边界正确；取消再打开无旧值；无效输入不提交"],
    ["异步行为", "区分加载、空、失败；避免旧响应覆盖与重复提交"],
    ["权限", "无权限或只读行为正确；前端控制与后端鉴权分别核对"],
    ["视觉与可用性", "同视口对照设计；长文本、滚动、键盘与操作入口可用"],
    ["真实接口", "Mock 验证与真实联调分别报告，不把模拟成功视为生产成功"]
], [3.4, 13.1])
h("交付应包含什么")
bullets([
    "源项目、页面、路由和可获取的设计 revision。",
    "实际修改文件、使用的依赖与公共组件、接口接入和模拟数据范围。",
    "实际执行的命令、通过或失败结果，以及真实入口的交互与视觉验证。",
    "未解决的接口、权限、资产或运行环境限制。"
])
p("代码生成完成后再进入团队的评审、提交和上线流程。设计截图不证明交互已经实现；文件已生成也不证明应用已经可运行。")

page("15 常见问题与排查")
h("账号和项目")
bullets([
    "登录后看不到团队项目：先核对所用账号与登录方式，再核对组织、工作区和邀请权限；注册账号不会自动加入既有团队。",
    "注册失败或收不到邮件：核对用户名、密码规则和邮箱；检查垃圾邮件。页面提示自助注册未开启时联系管理员。",
    "忘记密码：从登录页进入“无法登录？”，按当前支持流程处理；不能假定自助重置已开放。",
    "浏览器授权成功但 Desktop 未登录：回到授权页面打开 Desktop，保持应用运行并重试登录。"
])
h("AI 与 MCP")
bullets([
    "AI 看不到 Plasmic 工具：检查 Desktop 是否运行且已登录，MCP 开关是否开启；重启客户端或刷新配置。",
    "AI 连接了错误项目：先让 AI 读取应用状态，确认项目名称与链接；开始编辑前纠正目标。",
    "Skill 未安装到目标客户端：先用 skill install --dry-run 查看检测结果，安装或初始化客户端后重试 skill install。",
    "移动应用后连接失败：重新打开 Desktop 的 MCP 窗口并生成配置，避免使用旧的绝对路径。"
])
h("CLI 和资源")
code("plasmickit --version\nplasmickit version check\nplasmickit references check\nplasmickit skill install --dry-run")
bullets([
    "找不到 plasmickit 命令：确认 Node/npm 可用，重新打开终端；也可使用 npx -y @plasmickit/cli@latest。",
    "切换 Node 版本后 CLI 变旧：全局 npm 包通常按 Node 安装目录隔离，在当前环境重新安装最新版。",
    "资源加载失败：检查网络、代理和资源源可达性，修复后重跑 context resolve；不要把失败当成已加载最新规范。",
    "生成页面只有外观：补齐事件、状态与接口契约，并在 Preview 或实际工程中验证。"
])
p("排查时提供客户端与应用版本、目标项目、操作步骤和具体错误提示。不要在反馈或文档中附密码、验证码或会话凭据。")

page("16 入口  版本与截图说明")
h("常用入口")
link("产品与 Desktop 下载", "https://shiguanglab.com/projects/zhijie#download")
link("Studio 网页编辑器", "https://studio.plasmic.shiguanglab.com/")
link("拾光账号登录", "https://shiguanglab.com/login")
link("拾光账号注册", "https://shiguanglab.com/register")
link("登录帮助", "https://shiguanglab.com/login-help")
link("Node.js 下载", "https://nodejs.org/en/download")
link("Plasmickit npm 包", "https://www.npmjs.com/package/@plasmickit/cli")
h("本次核对的版本")
table(["产物", "核对版本", "更新入口"], [
    ["Studio", "0.0.85", "在线服务由维护者更新"],
    ["Desktop", "0.0.62", "官网下载或应用软件更新"],
    ["Plasmickit CLI", "0.0.84", "npm install -g @plasmickit/cli@latest"],
    ["规范资源", "0.0.87", "context resolve 或 references update"]
], [3.5, 3.1, 9.9])
p("下载时以官网最新可用版本为准。版本号是本手册核对记录，不是必须固定使用的版本。示例命令使用 @latest。")
h("界面与示例来源")
p("下载、注册、网页登录、项目创建、编辑器、AI、预览和分享图于 2026 年 10 月 14 日从实际页面截取。Desktop 登录图沿用仓库 2026 年 10 月 11 日界面记录；MCP 图沿用 2026 年 10 月 10 日实现记录，其入口与客户端列表已按当前源码核对。")
p("编辑与预览图来自 Admin Templates 素材库，使用示例数据。分享图只保留操作区，未展示个人邮箱。截图时未创建新账号、提交 AI 请求、发布项目或改变共享权限。")
p("操作依据还包括当前 Desktop README、CLI 实现、注册页面实现，以及 CLI 加载的 prototype／codegen 工作流、模型读取和验收规范。不同 UI 语言、窗口尺寸和后续版本可能改变标签与排版，操作以当前可见界面为准。")

# Remove template borders and theme font overrides from all styles and runs.
for element in [doc.styles.element, doc.element]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for fonts in element.xpath('.//w:rFonts'):
        for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
            fonts.attrib.pop(qn("w:" + attr), None)
        fonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
doc.save(OUT)
print(OUT)
